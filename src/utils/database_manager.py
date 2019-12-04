import openpyxl
from openpyxl.chart import PieChart, Reference
from os import path
import atexit


class DataBaseManager:
    def __init__(self, db_name):
        self.db_name = db_name

        if path.exists(self.db_name):
            self.wb = openpyxl.load_workbook(self.db_name)
            self.sheet = self.wb.active
        else:
            self.wb = openpyxl.Workbook()
            self.sheet = self.wb.active

            data = Reference(self.sheet, min_col=4, min_row=1, max_col=6, max_row=2)

            chart = PieChart()
            chart.title = 'Статистика'
            chart.add_data(data, titles_from_data=True)

            self.sheet.add_chart(chart, 'D4')

            self.sheet['D1'] = 'Негативных'
            self.sheet['D2'] = '=COUNTIF(B3:B20000,"Негативный")'

            self.sheet['E1'] = 'Нейтральных'
            self.sheet['E2'] = '=COUNTIF(B3:B20000,"Нейтральный")'

            self.sheet['F1'] = 'Позитивных'
            self.sheet['F2'] = '=COUNTIF(B3:B20000,"Позитивный")'

        atexit.register(self.__save)

    def append_result(self, result):
        self.sheet.append(result)

    def __save(self):
        self.wb.save(self.db_name)

